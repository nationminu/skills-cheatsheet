#!/usr/bin/env python3
"""PPTX에서 개수 세기 표현(PPT-MASTER.md 4.4) 찾기

사용:
  python3 scripts/check-phrases.py 결과.pptx          # 발견 시 종료 코드 1
  python3 scripts/check-phrases.py 결과.pptx --all    # 본문 불릿까지 전부 (기본은 제목·부제·블록 제목 추정 위치)
  python3 scripts/check-phrases.py --text "변수 아홉 개의 구성"   # 문구 하나 확인

- 찾는 형태: {숫자|한글 수사} + 가지·개·건·장면·단계·절·줄·묶음·항목·종
- 예외(사실 숫자: 날짜, 시각, 사양, 기간, 금액, KPI 값)는 사람이 결과를 보고 판단 — 이 스크립트는 후보만 보여 줌
- python-pptx 필요: pip install python-pptx
"""
import re
import sys

NUM = r"(?:\d+|한|두|세|네|다섯|여섯|일곱|여덟|아홉|열(?:\s?(?:한|두|세|네|다섯|여섯|일곱|여덟|아홉))?|스물|스무)"
UNIT = r"(?:가지|개|건|장면|단계|절|줄|묶음|항목|종)"
PATTERN = re.compile(rf"(?<![가-힣0-9]){NUM}\s?{UNIT}(?=$|[^가-힣]|의|와|과|을|를|은|는|이|가|로|으로|및|에|씩|만|도)")

# 오탐을 줄이기 위한 예외 (단어 일부로 쓰인 경우)
IGNORE = re.compile(r"(한\s?줄\s?요약|한\s?개념|세\s?가지\s?이상)")


def find(text):
    hits = []
    for m in PATTERN.finditer(text):
        if IGNORE.search(text[max(0, m.start() - 2): m.end() + 4]):
            continue
        hits.append(m.group(0))
    return hits


def is_heading(shape, para):
    """제목·부제·블록 제목 추정: 제목 자리표시자, 14 pt 이상 또는 굵은 한 줄"""
    if shape.is_placeholder and shape.placeholder_format.type is not None:
        if "TITLE" in str(shape.placeholder_format.type):
            return True
    runs = [r for r in para.runs if r.text.strip()]
    if not runs:
        return False
    sizes = [r.font.size.pt for r in runs if r.font.size is not None]
    bold = all(r.font.bold for r in runs)
    return (sizes and max(sizes) >= 14) or bold


def scan(path, all_text):
    from pptx import Presentation

    prs = Presentation(path)
    found = 0
    for no, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            frames = []
            if shape.has_text_frame:
                frames.append(shape.text_frame)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        frames.append(cell.text_frame)
            for tf in frames:
                for para in tf.paragraphs:
                    text = "".join(r.text for r in para.runs).strip()
                    if not text:
                        continue
                    if not all_text and shape.has_text_frame and not is_heading(shape, para):
                        continue
                    hits = find(text)
                    if hits:
                        found += 1
                        print(f"[{no:02d}장] {', '.join(hits)}  ←  {text}")
    if found:
        print(f"\n개수 세기 표현 후보 {found}곳 — 사실 숫자·KPI가 아니면 PPT-MASTER.md 4.4 표대로 수정")
    else:
        print("개수 세기 표현 없음")
    return found


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    if sys.argv[1] == "--text":
        hits = find(" ".join(sys.argv[2:]))
        print(", ".join(hits) if hits else "없음")
        sys.exit(1 if hits else 0)
    sys.exit(1 if scan(sys.argv[1], "--all" in sys.argv) else 0)


if __name__ == "__main__":
    main()
